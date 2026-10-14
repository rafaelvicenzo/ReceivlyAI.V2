"""Geração de laudo em Word (.docx) a partir do template da imobiliária."""

from __future__ import annotations

import io
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt
from docxtpl import DocxTemplate, InlineImage

from app import calculo_avaliacao, models, schemas
from app.calculo_avaliacao import FATORES_KEYS

BASE_DIR = Path(__file__).resolve().parent.parent
UPLOADS_DIR = BASE_DIR / "uploads"


def _fmt_brl(valor: float) -> str:
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fmt_num(valor: float | None, casas: int = 2) -> str:
    if valor is None:
        return "—"
    return f"{valor:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _data_longa(data_iso: str) -> str:
    try:
        dt = datetime.fromisoformat(data_iso[:10])
        meses = [
            "janeiro", "fevereiro", "março", "abril", "maio", "junho",
            "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
        ]
        return f"{dt.day} de {meses[dt.month - 1]} de {dt.year}"
    except ValueError:
        return data_iso


def _numero_por_extenso(n: int) -> str:
    un = [
        "zero", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito", "nove", "dez",
        "onze", "doze", "treze", "catorze", "quinze", "dezesseis", "dezessete", "dezoito", "dezenove",
    ]
    dz = ["", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta", "setenta", "oitenta", "noventa"]
    ct = [
        "", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos",
        "seiscentos", "setecentos", "oitocentos", "novecentos",
    ]

    def ate999(x: int) -> str:
        if x == 100:
            return "cem"
        c = x // 100
        r = x % 100
        partes = []
        if c:
            partes.append(ct[c])
        if r:
            if r < 20:
                partes.append(un[r])
            else:
                d, u = divmod(r, 10)
                partes.append(f"{dz[d]} e {un[u]}" if u else dz[d])
        return " e ".join(partes)

    if n == 0:
        return "zero"
    milhoes = n // 1_000_000
    milhares = (n % 1_000_000) // 1000
    resto = n % 1000
    grupos = []
    if milhoes:
        grupos.append(f"{ate999(milhoes)} {'milhão' if milhoes == 1 else 'milhões'}")
    if milhares:
        grupos.append("mil" if milhares == 1 else f"{ate999(milhares)} mil")
    if resto:
        grupos.append(ate999(resto))
    return " e ".join(grupos)


def _paginas_por_extenso(num_paginas: int | None) -> str:
    if not num_paginas or num_paginas < 1:
        return "[número por extenso]"
    return _numero_por_extenso(num_paginas)


def _titulo_laudo_capa(laudo: schemas.LaudoExportRequest) -> str:
    natureza = "RURAL" if (laudo.natureza_imovel or "urbano") == "rural" else "URBANO"
    if laudo.finalidade == "locacao":
        return f"LAUDO TÉCNICO DE AVALIAÇÃO DE LOCAÇÃO DE IMÓVEL {natureza}"
    return f"LAUDO TÉCNICO DE AVALIAÇÃO MERCADOLÓGICA DE IMÓVEL {natureza}"


def _nome_profissional(laudo: schemas.LaudoExportRequest) -> str:
    return (laudo.nome_profissional or laudo.perita or "").strip()


def _qualificacao_laudo(laudo: schemas.LaudoExportRequest) -> str:
    if laudo.qualificacao_laudo and laudo.qualificacao_laudo.strip():
        return laudo.qualificacao_laudo.strip()
    partes = [p for p in (laudo.creci, laudo.cnai) if p and p.strip()]
    if not partes:
        return "[informe CRECI/CNAI em Configurações]"
    return (
        f"{', '.join(partes)}, conforme Resolução COFECI n.º 1066/2007 e "
        "Ato Normativo COFECI n.º 001/2011 para emitir Parecer Técnico de Avaliação Mercadológica."
    )


def _endereco_completo(av: models.Avaliacao) -> str:
    partes = [av.endereco or "", av.numero or ""]
    if av.complemento:
        partes.append(av.complemento)
    loc = ", ".join(p for p in [av.bairro, f"{av.cidade}/{av.estado}" if av.cidade else av.estado] if p)
    if loc:
        partes.append(loc)
    return ", ".join(p for p in partes if p)


_ESTADOS_NOME = {
    "AC": "Acre", "AL": "Alagoas", "AP": "Amapá", "AM": "Amazonas", "BA": "Bahia", "CE": "Ceará",
    "DF": "Distrito Federal", "ES": "Espírito Santo", "GO": "Goiás", "MA": "Maranhão",
    "MT": "Mato Grosso", "MS": "Mato Grosso do Sul", "MG": "Minas Gerais", "PA": "Pará",
    "PB": "Paraíba", "PR": "Paraná", "PE": "Pernambuco", "PI": "Piauí", "RJ": "Rio de Janeiro",
    "RN": "Rio Grande do Norte", "RS": "Rio Grande do Sul", "RO": "Rondônia", "RR": "Roraima",
    "SC": "Santa Catarina", "SP": "São Paulo", "SE": "Sergipe", "TO": "Tocantins",
}

_DEFAULT_TEXTO_LOCALIZACAO = (
    "O imóvel avaliando ({{tipoImovel}}) está localizado na {{endereco}}, n.º {{numero}}"
    "{{complementoSuffix}}, no bairro de {{bairro}}, no município de {{cidade}}, no Estado de {{estadoNome}}."
)

_DEFAULT_TEXTO_VALORES_RESULTANTES = (
    "Os valores resultantes encontram-se dentro da média pesquisada, quando considerados todos "
    "os fatores que influenciam na sua formação."
)

_DEFAULT_TEXTO_ENCERRAMENTO = (
    "O presente laudo técnico de avaliação é composto de {{numPaginas}} ({{numPaginasExtenso}}) páginas, "
    "todas de um só lado, e assinado pelo profissional que subscreve esta última."
)

_DEFAULT_METODOLOGIA_COMPARATIVO = (
    "Para determinação do {{valorTipo}}, utilizou-se o Método Comparativo Direto de Dados. Tal método "
    "estabelece de forma direta um sistema de comparação com imóveis similares, através da coleta de "
    "dados em pesquisa junto a imobiliárias, profissionais especializados, proprietários, anúncios "
    "classificados, órgãos públicos e outros veículos que possam fornecer informações básicas e "
    "concretas dos imóveis ofertados ou vendidos recentemente.\n\n"
    "A similaridade dos imóveis pesquisados é de fundamental importância para que tenhamos elementos "
    "que possam ser comparados de forma direta.\n\n"
    "O processo comparativo, muito embora a pesquisa abranja imóveis similares em tamanho, padrão e "
    "local, requer uma melhoria no seu grau de precisão, obtida por outro processo auxiliar, "
    "denominado Homogeneização de Valores. Tal processo, recomendado por normas técnicas, visa corrigir "
    "as discrepâncias dos elementos comparativos, através de fatores cuja aplicação de um ou outro "
    "depende de cada caso. Assim, os fatores fundamentais, como: localização, atualização, fonte, "
    "padrão construtivo, idade e conservação, etc., são levados em conta para que os elementos sejam "
    "mais comparáveis, não obstante a similaridade já pré-selecionada na pesquisa.\n\n"
    "O método comparativo, dado as características do imóvel em questão e pela viabilidade de sua "
    "aplicação ao caso, é o que mais se ajusta para a determinação do {{valorTipo}}.\n\n"
    "Para melhorar ainda mais a precisão, após o processo de homogeneização, a média obtida passa por "
    "processo de saneamento de valores que eventualmente situam-se muito distantes do seu valor médio, "
    "distorcendo o valor unitário. Em geral, utiliza-se um intervalo de confiança, eliminando-se os "
    "valores fora deste intervalo e recalculando-se a nova média (média saneada). Assim, o valor "
    "unitário obtido pode ser considerado como o real {{valorTipo}}.\n\n"
    "Justificativa da metodologia: dentro da região de influência na qual se insere o imóvel foi possível "
    "identificar uma amostragem em tamanho suficiente para se obter um valor médio confiável. Tal fato "
    "nos levou à aplicação do método comparativo, com maior precisão para o caso em questão."
)

_DEFAULT_METODOLOGIA_EVOLUTIVO = (
    "Para determinação do {{valorTipo}}, utilizou-se o Método Evolutivo, conforme preconiza a NBR "
    "14.653 e as práticas consolidadas pelo IBAPE.\n\n"
    "O método evolutivo determina o valor do imóvel pela soma do custo de reprodução ou reposição da "
    "edificação, depreciada por idade, estado de conservação e obsolescências, acrescido do valor do "
    "terreno, quando for o caso. A depreciação física foi considerada mediante o método Ross-Heidecke, "
    "tendo em vista a idade aparente e o estado de conservação verificados em vistoria.\n\n"
    "O custo unitário básico de construção foi obtido por meio de tabelas e boletins técnicos de "
    "referência (como CUB/SINAPI ou equivalente regional), aplicados ao padrão construtivo e às "
    "características do imóvel avaliando.\n\n"
    "Justificativa da metodologia: diante das peculiaridades do imóvel avaliando e da necessidade de "
    "refletir a formação de custo do bem edificado, optou-se pelo método evolutivo, por ser o mais "
    "adequado à determinação do {{valorTipo}} no presente caso."
)


def _default_metodologia(metodo: str | None) -> str:
    if (metodo or "").strip().lower() == "evolutivo":
        return _DEFAULT_METODOLOGIA_EVOLUTIVO
    return _DEFAULT_METODOLOGIA_COMPARATIVO


def _nome_estado(uf: str | None) -> str:
    chave = (uf or "").strip().upper()
    return _ESTADOS_NOME.get(chave, uf or "[estado]")


def _substituir_variaveis(texto: str, vars_map: dict[str, str]) -> str:
    out = texto or ""
    for chave, valor in vars_map.items():
        out = out.replace(f"{{{{{chave}}}}}", valor)
    return out


def _template_tem_marcadores(caminho: Path) -> bool:
    try:
        with zipfile.ZipFile(caminho) as zf:
            for nome in zf.namelist():
                if nome.startswith("word/") and nome.endswith(".xml"):
                    if b"{{" in zf.read(nome):
                        return True
    except (zipfile.BadZipFile, OSError):
        return False
    return False


def _montar_comparaveis(avaliacao: models.Avaliacao, resultado: schemas.ResultadoAvaliacaoResponse) -> list[dict]:
    linhas = []
    for i, calc in enumerate(resultado.calculos, start=1):
        c = calc.comparavel
        fat = calc.fatores
        linhas.append({
            "n": i,
            "endereco": c.endereco,
            "bairro": c.bairro or "",
            "area": _fmt_num(c.area, 0),
            "valor": _fmt_brl(c.valor),
            "valor_unitario": _fmt_brl(calc.valor_unitario_base),
            "valor_homogeneizado": _fmt_brl(calc.valor_unitario_homogeneizado),
            "fator_local": _fmt_num(fat.fator_local, 3),
            "fator_area": _fmt_num(fat.fator_area, 3),
            "fator_depreciacao": _fmt_num(fat.fator_depreciacao, 3),
            "fator_padrao": _fmt_num(fat.fator_padrao, 3),
            "fator_testada": _fmt_num(fat.fator_testada, 3),
            "fator_total": _fmt_num(fat.fator_total, 3),
        })
    return linhas


def montar_contexto(
    avaliacao: models.Avaliacao,
    imobiliaria: models.Imobiliaria,
    laudo: schemas.LaudoExportRequest,
    resultado: schemas.ResultadoAvaliacaoResponse,
    laudo_calc: dict | None = None,
) -> dict:
    eh_locacao = laudo.finalidade == "locacao"
    laudo_calc = laudo_calc or calculo_avaliacao.calcular_laudo_detalhado(avaliacao)
    valor_calc = laudo_calc["valor_calculado"]
    valor_final = laudo.valor_adotado if laudo.valor_adotado else round(valor_calc / 10) * 10
    complemento = f", {avaliacao.complemento}" if avaliacao.complemento else ""
    imovel_desc = (
        f"imóvel ({(avaliacao.tipo or '').lower()}) situado em {avaliacao.endereco}, "
        f"nº {avaliacao.numero}{complemento}, {avaliacao.bairro}, {avaliacao.cidade}/{avaliacao.estado}"
    )

    complemento_suffix = f", {avaliacao.complemento}" if avaliacao.complemento else ""
    vars_texto = {
        "solicitante": laudo.solicitante or "[solicitante]",
        "objetivo": "determinar o valor de locação mensal" if eh_locacao else "determinar o valor de mercado",
        "valorTipo": "valor de locação mensal" if eh_locacao else "valor de mercado",
        "parteA": "locatário" if eh_locacao else "comprador",
        "parteB": "locador" if eh_locacao else "vendedor",
        "negocio": "a locação" if eh_locacao else "a venda",
        "imovel": imovel_desc,
        "avaliador": _nome_profissional(laudo),
        "profissional": _nome_profissional(laudo),
        "tituloProfissional": laudo.titulo_profissional or "",
        "qualificacao": _qualificacao_laudo(laudo),
        "creci": laudo.creci,
        "cnai": laudo.cnai,
        "matricula": laudo.matricula or avaliacao.matricula or "",
        "tipoImovel": (avaliacao.tipo or "imóvel").lower(),
        "endereco": avaliacao.endereco or "[endereço]",
        "numero": avaliacao.numero or "s/n",
        "complementoSuffix": complemento_suffix,
        "bairro": avaliacao.bairro or "[bairro]",
        "cidade": avaliacao.cidade or "[cidade]",
        "estadoNome": _nome_estado(avaliacao.estado),
        "numPaginas": str(laudo.num_paginas or "[nº de páginas]"),
        "numPaginasExtenso": _paginas_por_extenso(laudo.num_paginas),
        "dataLonga": _data_longa(laudo.data),
        "local": laudo.local or "[local]",
    }

    textos_raw = dict(laudo.textos or {})
    if not (textos_raw.get("localizacao") or "").strip():
        textos_raw["localizacao"] = _DEFAULT_TEXTO_LOCALIZACAO
    if not (textos_raw.get("valoresResultantes") or "").strip():
        textos_raw["valoresResultantes"] = _DEFAULT_TEXTO_VALORES_RESULTANTES
    if not (textos_raw.get("metodologia") or "").strip():
        textos_raw["metodologia"] = _default_metodologia(laudo.metodo_avaliacao)
    if not (textos_raw.get("encerramento") or "").strip():
        textos_raw["encerramento"] = _DEFAULT_TEXTO_ENCERRAMENTO
    textos = {chave: _substituir_variaveis(valor, vars_texto) for chave, valor in textos_raw.items()}

    matricula = laudo.matricula or avaliacao.matricula or ""
    titulo_capa = _titulo_laudo_capa(laudo)

    ctx = {
        "solicitante": laudo.solicitante,
        "processo": laudo.processo or "",
        "matricula": matricula,
        "natureza_imovel": laudo.natureza_imovel or "urbano",
        "titulo_laudo": titulo_capa,
        "perita": _nome_profissional(laudo),
        "avaliador": _nome_profissional(laudo),
        "profissional": _nome_profissional(laudo),
        "nome_profissional": _nome_profissional(laudo),
        "titulo_profissional": laudo.titulo_profissional or "",
        "qualificacao_laudo": _qualificacao_laudo(laudo),
        "creci": laudo.creci,
        "cnai": laudo.cnai,
        "local": laudo.local,
        "data": laudo.data,
        "data_longa": _data_longa(laudo.data),
        "finalidade": laudo.finalidade,
        "valor_tipo": vars_texto["valorTipo"],
        "objetivo": vars_texto["objetivo"],
        "endereco": avaliacao.endereco or "",
        "numero": avaliacao.numero or "",
        "complemento": avaliacao.complemento or "",
        "bairro": avaliacao.bairro or "",
        "cidade": avaliacao.cidade or "",
        "estado": avaliacao.estado or "",
        "cep": avaliacao.cep or "",
        "tipo": avaliacao.tipo or "",
        "area_imovel": _fmt_num(avaliacao.area_imovel or 0, 0),
        "quartos": str(avaliacao.quartos or ""),
        "banheiros": str(avaliacao.banheiros or ""),
        "vagas": str(avaliacao.vagas or ""),
        "endereco_completo": _endereco_completo(avaliacao),
        "imovel": imovel_desc,
        "valor_estimado": valor_final,
        "valor_estimado_fmt": _fmt_brl(valor_final),
        "valor_extenso": laudo.valor_extenso or _fmt_brl(valor_final),
        "valor_unitario_fmt": _fmt_brl(laudo_calc["unitario_final"]),
        "laudo_calc": laudo_calc,
        "imobiliaria_nome": imobiliaria.nome,
        "cor_destaque_laudo": (imobiliaria.cor_destaque_laudo or "#F59E0B").upper(),
        "texto_introducao": textos.get("introducao", ""),
        "texto_premissas": textos.get("premissas", ""),
        "texto_resultados": textos.get("resultados", ""),
        "texto_localizacao": textos.get("localizacao", ""),
        "texto_mercado_atual": textos.get("mercadoAtual", ""),
        "texto_perspectivas": textos.get("perspectivas", ""),
        "texto_valores_resultantes": textos.get("valoresResultantes", ""),
        "texto_metodologia": textos.get("metodologia", ""),
        "texto_homogeneizacao": textos.get("homogeneizacao", ""),
        "texto_tratamento": textos.get("tratamento", ""),
        "texto_especificacao": textos.get("especificacao", ""),
        "texto_encerramento": textos.get("encerramento", ""),
        "textos": textos,
        "comparaveis": _montar_comparaveis(avaliacao, resultado),
        "quantidade_comparaveis": len(resultado.calculos),
        "conteudo_laudo": "",
    }
    ctx["conteudo_laudo"] = _montar_texto_completo(ctx)
    return ctx


def _caminho_absoluto_upload(relativo: str | None) -> Path | None:
    if not relativo:
        return None
    rel = relativo.replace("\\", "/").lstrip("/")
    if rel.startswith("uploads/"):
        rel = rel[len("uploads/") :]
    return UPLOADS_DIR / rel


def _montar_texto_completo(ctx: dict) -> str:
    linhas = [
        f"Solicitante: {ctx['solicitante']}",
        f"Processo: {ctx['processo']}" if ctx.get("processo") else "",
        f"Matrícula: {ctx['matricula']}" if ctx.get("matricula") else "",
        f"Imóvel: {ctx['endereco_completo']}",
        f"Valor estimado: {ctx['valor_estimado_fmt']} ({ctx['valor_extenso']})",
        "",
    ]
    for titulo, chave in [
        ("A. Introdução", "texto_introducao"),
        ("B. Premissas e ressalvas", "texto_premissas"),
        ("C. Resultados", "texto_resultados"),
        ("E. Localização", "texto_localizacao"),
        ("F.1 Mercado atual", "texto_mercado_atual"),
        ("F.2 Perspectivas futuras", "texto_perspectivas"),
        ("F.3 Valores resultantes", "texto_valores_resultantes"),
        ("H. Metodologia", "texto_metodologia"),
        ("Homogeneização de valores", "texto_homogeneizacao"),
        ("I. Tratamento por fatores", "texto_tratamento"),
        ("N. Especificação da avaliação", "texto_especificacao"),
        ("O. Encerramento", "texto_encerramento"),
    ]:
        texto = ctx.get(chave, "")
        if texto and texto.strip():
            linhas.extend([titulo, texto, ""])
    linhas.append(f"{ctx['local']}, {ctx['data_longa']}.")
    linhas.extend([ctx["nome_profissional"], ctx["creci"], ctx["cnai"]])
    return "\n".join(l for l in linhas if l is not None)


def _adicionar_paragrafos(doc: Document, titulo: str, texto: str) -> None:
    if not texto or not texto.strip():
        return
    doc.add_heading(titulo, level=2)
    for bloco in texto.split("\n\n"):
        bloco = bloco.strip()
        if bloco:
            p = doc.add_paragraph(bloco)
            p.paragraph_format.space_after = Pt(8)
            for run in p.runs:
                run.font.size = Pt(11)


def _aplicar_estilo_tabela(tabela) -> None:
    for nome in ("Table Grid", "Normal Table", "Tabela com grade"):
        try:
            tabela.style = nome
            return
        except KeyError:
            continue


def _celula_texto(celula, texto: str, *, negrito: bool = False, italico: bool = False, tamanho: int = 8) -> None:
    celula.text = ""
    p = celula.paragraphs[0]
    run = p.add_run(texto)
    run.bold = negrito
    run.italic = italico
    run.font.size = Pt(tamanho)


def _adicionar_quadro_dados(doc: Document, titulo: str, linhas: list[tuple[str, str]]) -> None:
    doc.add_paragraph(titulo).runs[0].bold = True
    tabela = doc.add_table(rows=len(linhas), cols=2)
    _aplicar_estilo_tabela(tabela)
    for i, (rotulo, valor) in enumerate(linhas):
        _celula_texto(tabela.rows[i].cells[0], rotulo, negrito=True)
        _celula_texto(tabela.rows[i].cells[1], valor)
    doc.add_paragraph()


def _adicionar_tabela_homogeneizacao_k(doc: Document, laudo_calc: dict, *, eh_locacao: bool) -> None:
    if not laudo_calc.get("linhas"):
        return

    sufixo = " mês" if eh_locacao else ""
    san = laudo_calc["saneamento"]
    fatores_labels = ["Local", "Área", "Depreciação", "Padrão", "Testada"]

    doc.add_heading("K. Homogeneização", level=2)

    cols = 11
    tabela = doc.add_table(rows=2, cols=cols)
    _aplicar_estilo_tabela(tabela)

    hdr = tabela.rows[0].cells
    _celula_texto(hdr[0], "N.º", negrito=True)
    _celula_texto(hdr[1], f"Unitário (R$/m²{sufixo})", negrito=True)
    _celula_texto(hdr[2], "Fator oferta", negrito=True)
    _celula_texto(hdr[3], "Unitário base", negrito=True)
    _celula_texto(hdr[4], "Fator de ajuste", negrito=True)
    _celula_texto(hdr[9], "Unitário homog.", negrito=True)
    _celula_texto(hdr[10], "Unitário saneado", negrito=True)
    hdr[4].merge(hdr[8])

    sub = tabela.rows[1].cells
    for i, lbl in enumerate(fatores_labels):
        _celula_texto(sub[4 + i], lbl, negrito=True)

    row_av = tabela.add_row().cells
    _celula_texto(row_av[0], "Aval.", negrito=True, italico=True)
    for i in range(5):
        _celula_texto(row_av[4 + i], "1,000", italico=True)

    for linha in laudo_calc["linhas"]:
        row_fat = tabela.add_row().cells
        for i, chave in enumerate(FATORES_KEYS):
            _celula_texto(row_fat[4 + i], _fmt_num(linha["fatores"][chave], 3), italico=True)

        row = tabela.add_row().cells
        _celula_texto(row[0], str(linha["n"]))
        _celula_texto(row[1], _fmt_num(linha["unitario_bruto"]))
        _celula_texto(row[2], _fmt_num(linha["fator_oferta"]))
        _celula_texto(row[3], _fmt_num(linha["base"]))
        for i, chave in enumerate(FATORES_KEYS):
            _celula_texto(row[4 + i], _fmt_num(linha["ajustes"][chave]))
        _celula_texto(row[9], _fmt_num(linha["homogeneizado"]), negrito=True)
        saneado_txt = _fmt_num(linha["homogeneizado"]) if linha["saneado"] else "excluído"
        _celula_texto(row[10], saneado_txt, negrito=True)

    row_media = tabela.add_row().cells
    _celula_texto(row_media[0], "Média seca", negrito=True)
    _celula_texto(row_media[1], _fmt_num(laudo_calc["media_seca"]), negrito=True)
    _celula_texto(row_media[4], "Médias", negrito=True, italico=True)
    _celula_texto(row_media[9], _fmt_num(laudo_calc["media_homog"]), negrito=True)
    _celula_texto(row_media[10], _fmt_num(san["media_saneada"]), negrito=True)

    doc.add_paragraph()
    _adicionar_quadro_dados(doc, "HOMOGENEIZAÇÃO", [
        ("Número de elementos", str(len(laudo_calc["linhas"]))),
        (f"Unitário médio homogeneizado (R$/m²{sufixo})", _fmt_num(laudo_calc["media_homog"])),
        (f"Limite superior (R$/m²{sufixo}) (+30%)", _fmt_num(laudo_calc["limite_sup"])),
        (f"Limite inferior (R$/m²{sufixo}) (-30%)", _fmt_num(laudo_calc["limite_inf"])),
    ])
    _adicionar_quadro_dados(doc, "SANEAMENTO", [
        ("Número de elementos saneados", str(san["n"])),
        (f"Unitário médio saneado (R$/m²{sufixo})", _fmt_num(san["media_saneada"])),
        ("Desvio padrão", _fmt_num(san["desvio"])),
        ("Coeficiente de variação", _fmt_num(san["coef_variacao"], 4)),
        ("t de Student", _fmt_num(san["t"])),
        (f"Máximo (R$/m²{sufixo})", _fmt_num(san["maximo"])),
        (f"Mínimo (R$/m²{sufixo})", _fmt_num(san["minimo"])),
        ("Intervalo de confiabilidade", f"{_fmt_num(san['intervalo'])}%"),
    ])


def _adicionar_formacao_valor(doc: Document, laudo_calc: dict, ctx: dict, *, eh_locacao: bool) -> None:
    sufixo = " mês" if eh_locacao else ""
    titulo = "L. Cálculo do valor de locação" if eh_locacao else "L. Cálculo do valor de mercado"
    doc.add_heading(titulo, level=2)
    rotulo_valor = "VALOR DE LOCAÇÃO (R$/MÊS)" if eh_locacao else "VALOR DE MERCADO (R$)"
    _adicionar_quadro_dados(doc, f"FORMAÇÃO DO {ctx['valor_tipo'].upper()}", [
        ("ÁREA (m²)", _fmt_num(laudo_calc["area_imovel"], 0)),
        (f"VALOR UNITÁRIO FINAL (R$/m²{sufixo})", _fmt_num(laudo_calc["unitario_final"])),
        ("APROVEITAMENTO", "1,00"),
        (rotulo_valor, _fmt_brl(ctx["valor_estimado"])),
    ])


def _adicionar_tabela_comparaveis(doc: Document, comparaveis: list[dict]) -> None:
    if not comparaveis:
        return
    doc.add_heading("J.2. Fichas de pesquisa (resumo)", level=2)
    cols = ["Nº", "Endereço", "Área m²", "Valor", "Unit. bruto", "Unit. homog."]
    tabela = doc.add_table(rows=1, cols=len(cols))
    _aplicar_estilo_tabela(tabela)
    for i, col in enumerate(cols):
        tabela.rows[0].cells[i].text = col
    for c in comparaveis:
        row = tabela.add_row().cells
        row[0].text = str(c["n"])
        row[1].text = c["endereco"]
        row[2].text = c["area"]
        row[3].text = c["valor"]
        row[4].text = c.get("valor_unitario_bruto", c["valor_unitario"])
        row[5].text = c["valor_homogeneizado"]


def _preencher_corpo_laudo(
    doc: Document,
    ctx: dict,
    *,
    dentro_template: bool = False,
    logo_path: Path | None = None,
    laudo_calc: dict | None = None,
    omitir_cabecalho: bool = False,
) -> None:
    """Insere o laudo no corpo do documento. Em template, preserva cabeçalho/rodapé do Word."""
    if not dentro_template and not omitir_cabecalho:
        if logo_path and logo_path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run()
            run.add_picture(str(logo_path), width=Inches(2.2))
        doc.add_heading(ctx.get("titulo_laudo") or "Laudo de Avaliação", level=1)

    capa = doc.add_paragraph()
    capa.add_run(f"Solicitante: {ctx['solicitante']}\n").bold = True
    if ctx["processo"]:
        capa.add_run(f"Processo: {ctx['processo']}\n")
    if ctx.get("matricula"):
        capa.add_run(f"Matrícula: n.º {ctx['matricula']}\n")
    capa.add_run(f"Imóvel: {ctx['endereco_completo']}\n")
    capa.add_run(f"Valor estimado: {ctx['valor_estimado_fmt']} ({ctx['valor_extenso']})\n")

    secoes = [
        ("A. Introdução", ctx["texto_introducao"]),
        ("B. Premissas e ressalvas", ctx["texto_premissas"]),
        ("C. Resultados", ctx["texto_resultados"]),
        ("E. Localização", ctx["texto_localizacao"]),
        ("F.1 Mercado atual", ctx["texto_mercado_atual"]),
        ("F.2 Perspectivas futuras", ctx["texto_perspectivas"]),
        ("F.3 Valores resultantes", ctx["texto_valores_resultantes"]),
        ("H. Metodologia", ctx["texto_metodologia"]),
        ("Homogeneização de valores", ctx["texto_homogeneizacao"]),
        ("I. Tratamento por fatores", ctx["texto_tratamento"]),
        ("N. Especificação da avaliação", ctx["texto_especificacao"]),
        ("O. Encerramento", ctx["texto_encerramento"]),
    ]
    for titulo, texto in secoes:
        _adicionar_paragrafos(doc, titulo, texto)

    laudo_calc = laudo_calc or ctx.get("laudo_calc") or {}
    eh_locacao = ctx.get("finalidade") == "locacao"
    _adicionar_tabela_comparaveis(doc, ctx["comparaveis"])
    _adicionar_tabela_homogeneizacao_k(doc, laudo_calc, eh_locacao=eh_locacao)
    _adicionar_formacao_valor(doc, laudo_calc, ctx, eh_locacao=eh_locacao)

    assinatura = doc.add_paragraph()
    assinatura.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    assinatura.add_run(f"\n{ctx['local']}, {ctx['data_longa']}.\n\n").italic = True
    assinatura.add_run(f"{ctx['nome_profissional']}\n").bold = True
    assinatura.add_run(f"{ctx['creci']}\n")
    assinatura.add_run(ctx["cnai"])


def _mapa_substituicoes(ctx: dict) -> dict[str, str]:
    chaves = [
        "solicitante", "processo", "matricula", "natureza_imovel", "titulo_laudo", "perita", "avaliador", "profissional", "nome_profissional", "titulo_profissional", "qualificacao_laudo",
        "creci", "cnai", "local", "data", "data_longa",
        "finalidade", "valor_tipo", "objetivo", "endereco", "numero", "complemento",
        "bairro", "cidade", "estado", "cep", "tipo", "area_imovel", "endereco_completo",
        "imovel", "valor_estimado_fmt", "valor_extenso", "valor_unitario_fmt",
        "imobiliaria_nome", "texto_introducao", "texto_premissas", "texto_resultados",
        "texto_metodologia", "texto_homogeneizacao", "texto_tratamento",
        "texto_especificacao", "texto_encerramento", "conteudo_laudo",
    ]
    subs = {k: str(ctx.get(k, "") or "") for k in chaves}
    subs["VALOR_ESTIMADO"] = subs.get("valor_estimado_fmt", "")
    nome = subs.get("nome_profissional") or subs.get("avaliador") or subs.get("profissional") or subs.get("perita", "")
    subs["NOME_PROFISSIONAL"] = nome
    subs["PROFISSIONAL"] = nome
    subs["AVALIADOR"] = nome
    subs["AVALIADORA"] = nome
    subs["PERITA"] = nome
    subs["PERITO"] = nome
    subs["QUALIFICACAO"] = subs.get("qualificacao_laudo", "")
    subs["qualificacao"] = subs["QUALIFICACAO"]
    subs["tituloProfissional"] = subs.get("titulo_profissional", "")
    return subs


def _substituir_paragrafo(paragraph, subs: dict[str, str]) -> None:
    texto = paragraph.text
    if not texto:
        return
    novo = texto
    for chave, valor in subs.items():
        for padrao in (f"{{{{{chave}}}}}", f"{{{{ {chave} }}}}", f"[{chave.upper()}]", f"[{chave}]"):
            novo = novo.replace(padrao, valor)
    if novo != texto:
        paragraph.text = novo


def _substituir_marcadores_documento(doc: Document, ctx: dict) -> None:
    subs = _mapa_substituicoes(ctx)
    for paragrafo in doc.paragraphs:
        _substituir_paragrafo(paragrafo, subs)
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    _substituir_paragrafo(paragrafo, subs)
    for section in doc.sections:
        for bloco in (section.header, section.footer):
            if bloco is None:
                continue
            for paragrafo in bloco.paragraphs:
                _substituir_paragrafo(paragrafo, subs)


def _substituir_em_paragrafos(paragrafos, marcadores: tuple[str, ...], texto: str) -> bool:
    for paragrafo in paragrafos:
        for marcador in marcadores:
            if marcador in paragrafo.text:
                paragrafo.text = paragrafo.text.replace(marcador, texto)
                return True
    return False


def _substituir_placeholder_no_template(doc: Document, ctx: dict) -> bool:
    """Substitui marcador único no corpo, ex.: {{ conteudo_laudo }} ou [CONTEUDO]."""
    texto = _montar_texto_completo(ctx)
    marcadores = ("{{ conteudo_laudo }}", "{{conteudo_laudo}}", "[CONTEUDO]", "[CONTEUDO DO LAUDO]")

    if _substituir_em_paragrafos(doc.paragraphs, marcadores, texto):
        return True

    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                if _substituir_em_paragrafos(celula.paragraphs, marcadores, texto):
                    return True

    for section in doc.sections:
        for bloco in (section.header, section.footer):
            if bloco is None:
                continue
            if _substituir_em_paragrafos(bloco.paragraphs, marcadores, texto):
                return True

    return False


def _criar_laudo_padrao(ctx: dict, logo_path: Path | None, *, omitir_cabecalho: bool = False) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    _preencher_corpo_laudo(doc, ctx, dentro_template=False, logo_path=logo_path, omitir_cabecalho=omitir_cabecalho)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _renderizar_com_marcadores(template_path: Path, ctx: dict, logo_path: Path | None) -> bytes:
    tpl = DocxTemplate(str(template_path))
    render_ctx = {
        k: v for k, v in ctx.items()
        if k != "laudo_calc" and k != "textos" and not isinstance(v, dict)
    }
    render_ctx["comparaveis"] = ctx.get("comparaveis", [])
    if logo_path and logo_path.exists():
        render_ctx["logo"] = InlineImage(tpl, str(logo_path), width=Inches(2.2))
    tpl.render(render_ctx)
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def _preencher_template_docx(
    template_path: Path,
    ctx: dict,
    laudo_calc: dict,
    logo_path: Path | None,
) -> bytes:
    """Preenche .docx: marcadores Jinja, substituição de campos ou corpo completo."""
    if _template_tem_marcadores(template_path):
        try:
            return _renderizar_com_marcadores(template_path, ctx, logo_path)
        except Exception:
            pass

    doc = Document(str(template_path))
    _substituir_marcadores_documento(doc, ctx)
    placeholder_ok = _substituir_placeholder_no_template(doc, ctx)
    corpo_vazio = not any(p.text.strip() for p in doc.paragraphs)
    if not placeholder_ok or corpo_vazio:
        _preencher_corpo_laudo(doc, ctx, dentro_template=True, laudo_calc=laudo_calc, logo_path=logo_path)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _logo_path_para_laudo(imobiliaria: models.Imobiliaria) -> Path | None:
    if imobiliaria.ocultar_cabecalho_laudo or not imobiliaria.usar_marca_dagua:
        return None
    return _caminho_absoluto_upload(imobiliaria.marca_dagua_caminho)


def gerar_laudo_docx(
    avaliacao: models.Avaliacao,
    imobiliaria: models.Imobiliaria,
    laudo: schemas.LaudoExportRequest,
) -> tuple[bytes, str]:
    laudo_calc = calculo_avaliacao.calcular_laudo_detalhado(avaliacao)
    resultado = calculo_avaliacao.calcular_resultado(avaliacao)
    ctx = montar_contexto(avaliacao, imobiliaria, laudo, resultado, laudo_calc=laudo_calc)
    logo_path = _logo_path_para_laudo(imobiliaria)
    omitir_cab = bool(imobiliaria.ocultar_cabecalho_laudo)

    template_relativo = imobiliaria.template_laudo_caminho
    if not template_relativo:
        conteudo = _criar_laudo_padrao(ctx, logo_path, omitir_cabecalho=omitir_cab)
        nome = f"laudo_avaliacao_{avaliacao.id}.docx"
        return conteudo, nome

    template_path = _caminho_absoluto_upload(template_relativo)
    if not template_path or not template_path.exists():
        conteudo = _criar_laudo_padrao(ctx, logo_path, omitir_cabecalho=omitir_cab)
        return conteudo, f"laudo_avaliacao_{avaliacao.id}.docx"

    ext = template_path.suffix.lower()
    if ext == ".pdf":
        raise ValueError(
            "Templates PDF não podem ser preenchidos automaticamente. "
            "Envie um arquivo .docx com marcadores como {{ valor_estimado_fmt }}."
        )
    if ext == ".doc":
        raise ValueError("Arquivos .doc antigos não são suportados. Salve o template como .docx no Word.")

    conteudo = _preencher_template_docx(template_path, ctx, laudo_calc, logo_path)
    modo = "template" if _template_tem_marcadores(template_path) else "preenchido"

    base = Path(imobiliaria.template_laudo_nome or "laudo").stem
    return conteudo, f"{base}_{avaliacao.id}_{modo}.docx"


def _criar_laudo_pdf(ctx: dict, *, logo_path: Path | None = None, ocultar_cabecalho: bool = False) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    y = h - 50
    margem = 50
    largura = w - 2 * margem

    def linha(texto: str, *, negrito: bool = False, espaco: float = 14) -> None:
        nonlocal y
        if y < 60:
            c.showPage()
            y = h - 50
        c.setFont("Helvetica-Bold" if negrito else "Helvetica", 10)
        for parte in (texto or "").split("\n"):
            while len(parte) > 95:
                c.drawString(margem, y, parte[:95])
                parte = parte[95:]
                y -= espaco
                if y < 60:
                    c.showPage()
                    y = h - 50
            c.drawString(margem, y, parte)
            y -= espaco

    if not ocultar_cabecalho:
        if logo_path and logo_path.exists():
            try:
                from reportlab.lib.utils import ImageReader
                c.drawImage(ImageReader(str(logo_path)), margem, y - 40, width=120, height=40, preserveAspectRatio=True, mask="auto")
                y -= 55
            except Exception:
                pass
        linha(ctx.get("titulo_laudo") or "LAUDO TÉCNICO DE AVALIAÇÃO", negrito=True, espaco=18)
        linha("")
    for bloco in (ctx.get("conteudo_laudo") or "").split("\n"):
        linha(bloco.strip(), negrito=bloco.strip().endswith(":") and len(bloco) < 60)

    c.save()
    return buf.getvalue()


def _converter_docx_para_pdf(docx_bytes: bytes) -> bytes | None:
    """Converte DOCX preenchido em PDF (Word COM, docx2pdf ou LibreOffice)."""
    import subprocess
    import sys
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        docx_path = Path(tmp) / "laudo.docx"
        pdf_path = Path(tmp) / "laudo.pdf"
        docx_path.write_bytes(docx_bytes)

        if sys.platform == "win32":
            try:
                import win32com.client

                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                doc = word.Documents.Open(str(docx_path.resolve()))
                doc.SaveAs(str(pdf_path.resolve()), FileFormat=17)
                doc.Close(False)
                word.Quit()
                if pdf_path.exists():
                    return pdf_path.read_bytes()
            except Exception:
                try:
                    word.Quit()
                except Exception:
                    pass

        try:
            from docx2pdf import convert

            convert(str(docx_path), str(pdf_path))
            if pdf_path.exists():
                return pdf_path.read_bytes()
        except Exception:
            pass

        for cmd in ("soffice", "libreoffice"):
            try:
                subprocess.run(
                    [cmd, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx_path)],
                    check=True,
                    capture_output=True,
                    timeout=120,
                )
                if pdf_path.exists():
                    return pdf_path.read_bytes()
            except Exception:
                continue

    return None


def gerar_laudo_pdf(
    avaliacao: models.Avaliacao,
    imobiliaria: models.Imobiliaria,
    laudo: schemas.LaudoExportRequest,
) -> tuple[bytes, str]:
    """Gera PDF a partir do mesmo DOCX do template (Word/LibreOffice) ou layout simplificado."""
    laudo_calc = calculo_avaliacao.calcular_laudo_detalhado(avaliacao)
    resultado = calculo_avaliacao.calcular_resultado(avaliacao)
    ctx = montar_contexto(avaliacao, imobiliaria, laudo, resultado, laudo_calc=laudo_calc)
    logo_path = _logo_path_para_laudo(imobiliaria)
    omitir_cab = bool(imobiliaria.ocultar_cabecalho_laudo)
    base = Path(imobiliaria.template_laudo_nome or "laudo").stem
    nome_pdf = f"{base}_{avaliacao.id}.pdf"

    docx_bytes, _ = gerar_laudo_docx(avaliacao, imobiliaria, laudo)
    pdf_bytes = _converter_docx_para_pdf(docx_bytes)
    if pdf_bytes:
        return pdf_bytes, nome_pdf

    if imobiliaria.template_laudo_caminho:
        raise ValueError(
            "Não foi possível converter o laudo para PDF com seu template. "
            "No Windows, instale o Microsoft Word ou use «Gerar Word» e exporte para PDF pelo Word."
        )

    conteudo = _criar_laudo_pdf(ctx, logo_path=logo_path, ocultar_cabecalho=omitir_cab)
    return conteudo, nome_pdf
